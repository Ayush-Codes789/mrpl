"""
Deliverable generation — turns agent results into downloadable files.

Based on task type, produces:
  - coding  → .py file with the code
  - reasoning → .docx report (python-docx)
  - vision → .docx with extracted text/descriptions
  - general → .docx summary

Also supports .pptx and .xlsx when explicitly requested.
"""

import os
from datetime import datetime
from typing import List, Optional

import re
_OUTPUT_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "outputs"
)


def _ensure_output_dir(subdir: Optional[str] = None) -> str:
    d = os.path.join(_OUTPUT_DIR, subdir) if subdir else _OUTPUT_DIR
    os.makedirs(d, exist_ok=True)
    return d


def generate_code_file(code: str, filename: Optional[str] = None) -> str:
    """Save generated code to a .py file."""
    out_dir = _ensure_output_dir("code")
    if not filename:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"generated_{ts}.py"
    path = os.path.join(out_dir, filename)
    with open(path, "w", encoding="utf-8") as f:
        f.write(code)
    return path


def generate_docx(
    title: str,
    content: str,
    sources: Optional[List[str]] = None,
    filename: Optional[str] = None,
) -> str:
    """Generate a .docx report from agent results."""
    from docx import Document
    from docx.shared import Pt, Inches

    doc = Document()

    # Title
    doc.add_heading(title, level=1)

    # Metadata
    doc.add_paragraph(
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        style="Subtitle",
    )

    # Main content — split on double newlines for paragraphs
    doc.add_heading("Content", level=2)
    for para_text in content.split("\n\n"):
        para_text = para_text.strip()
        if para_text:
            doc.add_paragraph(para_text)

    # Sources
    if sources:
        doc.add_heading("Sources", level=2)
        for src in sources:
            doc.add_paragraph(src, style="List Bullet")

    out_dir = _ensure_output_dir("reports")
    if not filename:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"report_{ts}.docx"
    path = os.path.join(out_dir, filename)
    doc.save(path)
    return path


def generate_pptx(
    title: str,
    slides_content: List[dict],
    filename: Optional[str] = None,
) -> str:
    """Generate a .pptx presentation.

    slides_content: list of {"title": str, "body": str}
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()

    # Title slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = f"Generated {datetime.now().strftime('%Y-%m-%d')}"

    # Content slides
    content_layout = prs.slide_layouts[1]
    for sc in slides_content:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = sc.get("title", "")
        slide.placeholders[1].text = sc.get("body", "")

    out_dir = _ensure_output_dir("presentations")
    if not filename:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"presentation_{ts}.pptx"
    path = os.path.join(out_dir, filename)
    prs.save(path)
    return path


def generate_xlsx(
    title: str,
    data: dict,
    filename: Optional[str] = None,
) -> str:
    """Generate an .xlsx file from dict data.

    data: {"column_name": [values...], ...}
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]  # Excel limit

    # Headers
    headers = list(data.keys())
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)

    # Data rows
    max_rows = max(len(v) for v in data.values()) if data else 0
    for row in range(max_rows):
        for col, h in enumerate(headers, 1):
            vals = data[h]
            ws.cell(row=row + 2, column=col, value=vals[row] if row < len(vals) else "")

    out_dir = _ensure_output_dir("spreadsheets")
    if not filename:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"data_{ts}.xlsx"
    path = os.path.join(out_dir, filename)
    wb.save(path)
    return path


def generate_deliverable(task_type: str, content: str, title: Optional[str] = None,
                         sources: Optional[List[str]] = None) -> str:
    """High-level: pick the right format based on task_type and generate."""
    if not title:
        title = f"MRPL Workbench — {task_type.title()} Output"

    if task_type == "coding":
        return generate_code_file(content)
    else:
        return generate_docx(title, content, sources=sources)
